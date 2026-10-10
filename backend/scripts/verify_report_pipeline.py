"""报告链端到端验证（2026-09-28，修 G-06 / G-07 后）。

覆盖：建轮次 → 落产物 → 生成报告 → 下载文件 → 错误路径。
重点验三件事：
  1. 模板缺失时**生成占位报告**（真文件、可下载），status = placeholder；
  2. 下载端点能真取回文件（Content-Disposition 正确）；
  3. 错误路径如实：轮次不存在 404、没生成就下载 404。

跑法（后端 8011 在跑）：
    python scripts/verify_report_pipeline.py
"""
from __future__ import annotations

import json
import sys
import urllib.error
import urllib.parse
import urllib.request

BASE = "http://127.0.0.1:8011"
SCENE = "motor.cooling-fan.fortior-fu6832"
_failed = 0


def check(cond: bool, label: str, detail: str = "") -> bool:
    """断言并留痕。**必须返回 cond**（漏 return 会让调用方静默失效）。"""
    global _failed
    if cond:
        print(f"[OK]  {label}" + (f"  {detail}" if detail else ""))
    else:
        _failed += 1
        print(f"[NG]  {label}" + (f"  {detail}" if detail else ""))
    return cond


def req(method: str, path: str, body=None) -> tuple[int, dict, bytes]:
    data = json.dumps(body).encode("utf-8") if body is not None else None
    url = f"{BASE}{urllib.parse.quote(path, safe='/?&=%')}"
    r = urllib.request.Request(
        url, data=data, method=method, headers={"Content-Type": "application/json"}
    )
    try:
        with urllib.request.urlopen(r, timeout=180) as resp:
            raw = resp.read()
            ctype = resp.headers.get("Content-Type", "")
            disp = resp.headers.get("Content-Disposition", "")
            # ⚠️ 判「是不是文件下载」要看 **Content-Disposition: attachment**，
            # 不能只看 Content-Type——占位报告本身就是 JSON 文件，
            # 按 ctype 判会把它误判成普通 API 响应（首版就栽在这）。
            if disp or "application/json" not in ctype:
                return resp.status, {"_file": True, "_len": len(raw), "_disp": disp}, raw
            return resp.status, json.loads(raw.decode("utf-8")), raw
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw.decode("utf-8") or "{}"), raw
        except json.JSONDecodeError:
            return e.code, {}, raw


def main() -> int:
    print("== ① 建一轮并落 config-snapshot（模拟流程 2 的确认锁定）==")
    st, out, _ = req("POST", "/api/runs", {"scene_id": SCENE, "mode": "mock"})
    check(st == 200, "POST /api/runs", f"HTTP {st}")
    run_id = (out.get("data") or {}).get("run_id")
    check(bool(run_id), "拿到 run_id", str(run_id))
    if not run_id:
        return 1

    snap = {
        "schema": "config-snapshot/v1",
        "snapshot_id": f"cs-{run_id[:8]}",
        "scene_id": SCENE,
        "locked_at": "2026-09-28T12:00:00Z",
        "locked_by": "张工",
        "locked": True,
        "product_attributes": [
            {"key": "model", "value": "FU6832N", "origin": "document", "_class": "header"},
            {"key": "project_name", "value": "散热风扇启动测试", "origin": "manual", "_class": "header"},
            {"key": "firmware_version", "value": "v1.2.3", "origin": "document", "_class": "header"},
            {"key": "inlet_start_addr", "value": "0x1000", "origin": "manual", "_class": "header"},
            {"key": "bootloader", "value": "BL-2.0", "origin": "document", "_class": "header"},
            {"key": "outlet_start_addr", "value": "0x2000", "origin": "manual", "_class": "header"},
        ],
        "firmware": {"project_path": "", "param_header_path": ""},
        "interlock_instances": [],
        "channel_map": [],
    }
    st, out, _ = req("PUT", f"/api/runs/{run_id}/config-snapshot", {"content": snap})
    check(st == 200, "PUT config-snapshot", f"HTTP {st}")

    print("\n== ② 读回核对（报告页装配层读的就是这一条）==")
    st, out, _ = req("GET", f"/api/runs/{run_id}")
    check(st == 200, "GET /api/runs/{id}", f"HTTP {st}")
    arts = (out.get("data") or {}).get("artifacts") or []
    hit = next((a for a in arts if a.get("artifact_kind") == "config-snapshot"), None)
    check(hit is not None, "产物里有 config-snapshot")
    attrs = ((hit or {}).get("content") or {}).get("product_attributes") or []
    header_vals = {a.get("key"): a.get("value") for a in attrs if a.get("_class") == "header"}
    check(len(header_vals) == 6, "6 个报告表头字段可取回", f"{len(header_vals)} 项")
    check(header_vals.get("model") == "FU6832N", "model 值正确", str(header_vals.get("model")))

    print("\n== ③ 生成报告（模板缺失 → 应走占位档，且真产出文件）==")
    st, out, _ = req("POST", f"/api/reports/generate?run_id={run_id}", {})
    check(st == 200, "POST /api/reports/generate", f"HTTP {st}")
    d = out.get("data") or {}
    status = d.get("status")
    check(status in ("ok", "placeholder"), "状态是 ok 或 placeholder（不是 mock）", str(status))
    check(status == "placeholder", "本机无模板 → 走占位档", str(status))
    check(bool(d.get("file")), "返回了文件名", str(d.get("file")))
    check((d.get("bytes") or 0) > 0, "文件不是空的", f"{d.get('bytes')} 字节")
    # ⚠️ 文案与文件格式都改过，断言要跟着走，否则脚本永远报假 NG：
    #   2026-09-28 占位件从 JSON 改成 **xlsm**（能打开看效果）；2026-09-29 reason 措辞
    #   从「模板未到位」改为「使用系统自建模板（客户模板登记值：…）」。故：
    #   · reason 判「系统自建模板」；· 文件名判 .xlsm 且带 placeholder 标记。
    check("系统自建模板" in str(d.get("reason", "")), "reason 如实说明走系统自建模板")
    check(str(d.get("file", "")).endswith(".placeholder.xlsm"), "占位件是 xlsm 且带 placeholder 标记", str(d.get("file")))

    print("\n== ④ 下载报告文件（G-07）==")
    st, meta, raw = req("GET", f"/api/reports/{run_id}/download")
    check(st == 200, "GET /api/reports/{id}/download", f"HTTP {st}")
    check(meta.get("_file") is True, "带 Content-Disposition，是文件下载", str(meta.get("_disp"))[:60])
    check(len(raw) > 0, "下载到的内容非空", f"{len(raw)} 字节")
    # 占位件是 xlsm（二进制，ZIP 头 PK\x03\x04），不再是 JSON —— 断言改为「是合法 xlsx 容器」
    check(raw[:2] == b"PK", "下载到的文件是合法的 Excel（xlsx）容器", raw[:4].hex())

    print("\n== ⑤ 错误路径如实 ==")
    st, _, _ = req("POST", "/api/reports/generate?run_id=不存在的轮次", {})
    check(st == 404, "轮次不存在 → 404（不生成指向空轮次的报告）", f"HTTP {st}")
    st, _, _ = req("GET", "/api/reports/从未生成过的轮次/download")
    check(st == 404, "没生成就下载 → 404", f"HTTP {st}")

    print("\n== ⑥ 重复生成是覆盖而非堆叠 ==")
    st, out2, _ = req("POST", f"/api/reports/generate?run_id={run_id}", {})
    check(st == 200, "二次生成 HTTP 200", f"HTTP {st}")
    from pathlib import Path
    rd = Path(__file__).resolve().parent.parent / "uploads" / "reports"
    files = sorted(p.name for p in rd.glob(f"report-{run_id}.*")) if rd.is_dir() else []
    check(len(files) == 1, "同一轮只有 1 份报告文件（覆盖）", ",".join(files))

    # ── ⑦ 确认表头 → 生成报告：确认值必须真写进 xlsm（2026-09-30 新增）──
    #
    # 这是「确认表头是否真的传递到报告」的消费端断言：光看接口 200 不算数，
    # 必须**打开生成的 xlsm、读出表头区的值**，确认页面上确认过的值真的在里面。
    print("\n== ⑦ 确认表头写入 report-instance → 报告表头采用确认值 ==")
    confirmed_header = {
        # 报告页本流程补的 4 项里挑 2 项（config-snapshot 里**没有**，只能来自确认产物）
        "customer_name": "确-客户甲",
        "customer_engineer": "确-工程师乙",
        # 上游 6 项里挑 1 项，**值与快照不同**，用来验证「确认值优先于快照」
        "model": "FU6832-CONFIRMED",
    }
    fields = [
        {"key": "model", "name": "机型", "from_flow": "F2 初始化配置", "value_slot": "B3"},
        {"key": "customer_name", "name": "客户名称", "from_flow": "F6 报告", "value_slot": "B12"},
        {"key": "customer_engineer", "name": "客户工程师", "from_flow": "F6 报告", "value_slot": "B13"},
        {"key": "email", "name": "Email", "from_flow": "F6 报告", "value_slot": "B14"},
    ]
    inst = {
        "schema": "report-instance/v1",
        "scene_id": SCENE,
        "status": "confirmed",
        "header": confirmed_header,
        "header_total": len(fields),
        "fields": fields,
        "confirmed_at": "2026-09-30T09:00:00Z",
        "confirmed_by": "张工",
    }
    st, out, _ = req("PUT", f"/api/runs/{run_id}/report-instance", {"content": inst})
    check(st == 200, "PUT report-instance(confirmed)", f"HTTP {st}")

    st, out, _ = req("POST", f"/api/reports/generate?run_id={run_id}", {})
    check(st == 200, "确认后重新生成报告", f"HTTP {st}")

    # 打开生成的 xlsm，读 A/B 两列，构建「标签 → 值」映射
    xlsm = rd / f"report-{run_id}.placeholder.xlsm"
    check(xlsm.exists(), "占位 xlsm 文件存在", xlsm.name)
    if xlsm.exists():
        from openpyxl import load_workbook
        wb = load_workbook(str(xlsm))
        ws = wb.active
        pairs: dict[str, str] = {}
        for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
            k, v = (row + (None, None))[:2]
            if k is not None and v is not None:
                pairs[str(k).strip()] = str(v).strip()
        # 本流程补的 2 项必须来自确认产物（快照里根本没有这两个 key）
        check(pairs.get("客户名称") == "确-客户甲", "确认的「客户名称」真写进报告表头", str(pairs.get("客户名称")))
        check(
            pairs.get("客户工程师") == "确-工程师乙",
            "确认的「客户工程师」真写进报告表头",
            str(pairs.get("客户工程师")),
        )
        # 上游项：确认值覆盖快照原值（快照里 model=FU6832N）
        check(
            pairs.get("机型") == "FU6832-CONFIRMED",
            "确认值优先于 config-snapshot（model 被覆盖）",
            str(pairs.get("机型")),
        )
        # 未确认的字段如实标「未提供」，不编
        check(
            "未提供" in str(pairs.get("Email", "")),
            "未确认的字段如实显示「未提供」",
            str(pairs.get("Email")),
        )
        wb.close()

    print()
    if _failed:
        print(f"FAILED: {_failed} 项未通过")
        return 1
    print("ALL PASS")
    print(f"（测试轮次 {run_id} 与其报告文件保留，供复核）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
