"""临时脚本：把生成的占位报告 xlsm 内容打印出来（供人工核对效果）。

⚠️ 临时用途，核对完即删。之所以不写进正式 scripts/ 家族：
它只做「打印」不做事实验证，正式的验收在 `verify_report_xlsm.py`。

Windows 控制台默认 GBK，而占位横幅含 `⚠`（U+26A0）会直接
`UnicodeEncodeError` —— 这正是本项目记过的那条坑，所以这里显式
`reconfigure(encoding="utf-8")`。
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook  # noqa: E402

from app.core.paths import reports_dir  # noqa: E402


def main() -> int:
    run_id = sys.argv[1] if len(sys.argv) > 1 else "d5c08959241847798ab15884d0ae141a"
    p = reports_dir() / f"report-{run_id}.placeholder.xlsm"
    wb = load_workbook(str(p))
    for ws in wb.worksheets:
        print(f"=== 工作表：{ws.title}（{ws.max_row} 行 × {ws.max_column} 列）===")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                print(" | ".join(cells))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
